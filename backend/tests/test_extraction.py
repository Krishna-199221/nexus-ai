from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.services.extraction import extract_text


def test_extract_txt(tmp_path: Path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text(
        "NEXUS AI extraction test",
        encoding="utf-8",
    )

    result = extract_text(str(file_path), "TXT")

    assert result == "NEXUS AI extraction test"


def test_extract_unsupported_type(tmp_path: Path):
    file_path = tmp_path / "sample.pptx"
    file_path.write_bytes(b"not a real pptx")

    with pytest.raises(
        ValueError,
        match="Unsupported document type for extraction: PPTX",
    ):
        extract_text(str(file_path), "PPTX")


def test_extract_csv(tmp_path: Path):
    file_path = tmp_path / "resources.csv"
    file_path.write_text(
        "Name,Department,Utilization\n"
        "Rahul,Engineering,82\n"
        "Priya,Operations,67\n",
        encoding="utf-8",
    )

    result = extract_text(str(file_path), "CSV")

    assert result == (
        "Name | Department | Utilization\n"
        "Rahul | Engineering | 82\n"
        "Priya | Operations | 67"
    )


def test_extract_pdf(tmp_path: Path):
    file_path = tmp_path / "sample.pdf"

    import pymupdf

    document = pymupdf.open()

    page_one = document.new_page()
    page_one.insert_text((72, 72), "NEXUS AI page one")

    page_two = document.new_page()
    page_two.insert_text((72, 72), "NEXUS AI page two")

    document.save(file_path)
    document.close()

    result = extract_text(str(file_path), "PDF")

    assert result == (
        "[Page 1]\n"
        "NEXUS AI page one\n\n"
        "[Page 2]\n"
        "NEXUS AI page two"
    )


def test_extract_docx(tmp_path: Path):
    file_path = tmp_path / "sample.docx"

    document = DocxDocument()
    document.add_paragraph("NEXUS AI document extraction")
    document.add_paragraph("Resource intelligence platform")
    document.add_paragraph("")
    document.add_paragraph("AI-powered automation")

    document.save(file_path)

    result = extract_text(str(file_path), "DOCX")

    assert result == (
        "NEXUS AI document extraction\n"
        "Resource intelligence platform\n"
        "AI-powered automation"
    )


def test_extract_xlsx(tmp_path: Path):
    file_path = tmp_path / "resources.xlsx"

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Resources"

    worksheet.append(["Name", "Department", "Utilization"])
    worksheet.append(["Rahul", "Engineering", 82])
    worksheet.append(["Priya", "Operations", 67])

    second_sheet = workbook.create_sheet("Summary")
    second_sheet.append(["Metric", "Value"])
    second_sheet.append(["Total Resources", 2])

    workbook.save(file_path)
    workbook.close()

    result = extract_text(str(file_path), "XLSX")

    assert result == (
        "[Sheet: Resources]\n"
        "Name | Department | Utilization\n"
        "Rahul | Engineering | 82\n"
        "Priya | Operations | 67\n\n"
        "[Sheet: Summary]\n"
        "Metric | Value\n"
        "Total Resources | 2"
    )
