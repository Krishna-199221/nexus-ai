import csv
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from openpyxl import load_workbook


def extract_text(file_path: str, document_type: str) -> str:
    if document_type == "TXT":
        return Path(file_path).read_text(encoding="utf-8")

    if document_type == "CSV":
        return _extract_csv(file_path)

    if document_type == "PDF":
        return _extract_pdf(file_path)

    if document_type == "DOCX":
        return _extract_docx(file_path)

    if document_type == "XLSX":
        return _extract_xlsx(file_path)

    raise ValueError(
        f"Unsupported document type for extraction: {document_type}"
    )


def _extract_csv(file_path: str) -> str:
    rows = []

    with Path(file_path).open(
        "r",
        encoding="utf-8",
        newline="",
    ) as csv_file:
        reader = csv.reader(csv_file)

        for row in reader:
            rows.append(" | ".join(row))

    return "\n".join(rows)


def _extract_pdf(file_path: str) -> str:
    pages = []

    with pymupdf.open(file_path) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text().strip()

            if text:
                pages.append(
                    f"[Page {page_number}]\n{text}"
                )

    return "\n\n".join(pages)


def _extract_docx(file_path: str) -> str:
    document = DocxDocument(file_path)
    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    return "\n".join(paragraphs)


def _extract_xlsx(file_path: str) -> str:
    workbook = load_workbook(
        filename=file_path,
        read_only=True,
        data_only=True,
    )

    sheets = []

    try:
        for worksheet in workbook.worksheets:
            rows = []

            for row in worksheet.iter_rows(values_only=True):
                values = [
                    str(value).strip()
                    for value in row
                    if value is not None
                ]

                if values:
                    rows.append(" | ".join(values))

            if rows:
                sheets.append(
                    f"[Sheet: {worksheet.title}]\n"
                    + "\n".join(rows)
                )

        return "\n\n".join(sheets)
    finally:
        workbook.close()
